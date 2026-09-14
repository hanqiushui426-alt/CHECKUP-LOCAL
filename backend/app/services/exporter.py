"""openpyxl 导出 .xlsx：检验结果明细、单患者单指标趋势序列。"""
from __future__ import annotations

import io
import re
from typing import Optional

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import models_dao as dao
from .i18n import tr
from .template_store import normalize_text

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11, name="微软雅黑")
DATA_FONT = Font(name="微软雅黑", size=10)

FLAG_TEXT = {"high": {"zh-CN": "偏高", "en-US": "High"},
             "low": {"zh-CN": "偏低", "en-US": "Low"},
             "normal": {"zh-CN": "", "en-US": ""}}


def _flag_text(flag: str, lang: str) -> str:
    return FLAG_TEXT.get(flag, FLAG_TEXT["normal"]).get(lang, FLAG_TEXT[flag]["zh-CN"])
ABNORMAL_HIGH_FILL = PatternFill("solid", fgColor="FEF2F2")
ABNORMAL_LOW_FILL = PatternFill("solid", fgColor="E0F2F1")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_widths(ws, headers: list[str]) -> None:
    for idx, h in enumerate(headers, start=1):
        w = len(h) * 2
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            v = row[0].value
            if v is not None:
                s = str(v)
                w = max(w, sum(2 if ord(c) > 127 else 1 for c in s))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(10, w + 2), 42)


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"


def _style_body(ws) -> None:
    for row in ws.iter_rows(min_row=2, max_col=ws.max_column):
        for c in row:
            c.font = DATA_FONT
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=False)


def export_results_xlsx(patient_id: Optional[int] = None, item: Optional[str] = None,
                        report_type: Optional[str] = None, date_from: Optional[str] = None,
                        date_to: Optional[str] = None, lang: Optional[str] = None) -> bytes:
    rows = dao.export_rows(patient_id, item, report_type, date_from, date_to)
    lng = lang or "zh-CN"
    wb = Workbook()
    ws = wb.active
    ws.title = tr(lng, "sheet.results")
    headers = [tr(lng, "col.patient"), tr(lng, "col.gender"), tr(lng, "col.reportDate"),
               tr(lng, "col.testItem"), tr(lng, "col.hospital"), tr(lng, "col.item"),
               tr(lng, "col.value"), tr(lng, "col.valueNum"), tr(lng, "col.unit"),
               tr(lng, "col.ref"), tr(lng, "col.refLow"), tr(lng, "col.refHigh"),
               tr(lng, "col.flag"), tr(lng, "col.source")]
    _write_header(ws, headers)
    for r in rows:
        flag = _flag_text(r["flag"], lng)
        ws.append([
            r["patient_name"] or "", r["gender"] or "", r["report_date"] or "",
            r["report_type"] or "", r["hospital"] or "", r["item"] or "",
            r["value_text"] or "", r["value_num"] if r["value_num"] is not None else "",
            r["unit"] or "", r["ref_text"] or "",
            r["ref_low"] if r["ref_low"] is not None else "",
            r["ref_high"] if r["ref_high"] is not None else "",
            flag, r["source_filename"] or "",
        ])
        if r["flag"] == "high":
            fill = ABNORMAL_HIGH_FILL
        elif r["flag"] == "low":
            fill = ABNORMAL_LOW_FILL
        else:
            continue
        for cell in ws[ws.max_row]:
            cell.fill = fill
    _sheet_widths(ws, headers)
    _style_body(ws)
    return _to_bytes(wb)


def export_trend_xlsx(patient_id: int, item_norm: str, lang: Optional[str] = None) -> bytes:
    series = dao.trend_series(patient_id, item_norm)
    lng = lang or "zh-CN"
    wb = Workbook()
    ws = wb.active
    ws.title = tr(lng, "sheet.trend")
    headers = _trend_headers(lng)
    _write_header(ws, headers)
    for r in series:
        _append_trend_row(ws, r, lng)
    _sheet_widths(ws, headers)
    _style_body(ws)
    _hide_aux_columns(ws)
    name = (series[0].get("item") if series else None) or item_norm
    unit = next((r.get("unit") for r in series if r.get("unit")), "") or ""
    _add_trend_chart(ws, f"{name} · {tr(lng, 'sheet.trend')}", unit)
    return _to_bytes(wb)


# 趋势图配色（与页面一致）
_C_LINE = "0F766E"
_C_HIGH = "DC2626"
_C_LOW = "2563EB"
_C_REF = "F59E0B"
# 数据区列序（1 起）：报告日期 检验项目 报告类型 结果 数值 单位 参考下限 参考上限 异常 来源文件
_COL_VALUE = 5
_COL_REF_LOW = 7
_COL_REF_HIGH = 8
_COL_HIGH_PT = 11   # 辅助列：仅偏高点的数值
_COL_LOW_PT = 12    # 辅助列：仅偏低点的数值
_CHART_ANCHOR = "N2"


def _trend_headers(lng: str) -> list[str]:
    """趋势/批量导出统一的列（末两列是绘图辅助列，导出后自动隐藏）。"""
    return [tr(lng, "col.reportDate"), tr(lng, "col.testItem"), tr(lng, "col.reportType"),
            tr(lng, "col.value"), tr(lng, "col.valueNum"), tr(lng, "col.unit"),
            tr(lng, "col.refLow"), tr(lng, "col.refHigh"), tr(lng, "col.flag"),
            tr(lng, "col.source"), tr(lng, "col.highPoint"), tr(lng, "col.lowPoint")]


def _append_trend_row(ws, r: dict, lng: str) -> None:
    """写入一行趋势数据，并按异常标记整行着色。辅助列供图表标记异常点。"""
    flag = r.get("flag") or "normal"
    num = r.get("value_num")
    ws.append([
        r.get("report_date") or "",
        r.get("item") or "",
        r.get("report_type") or "",
        r.get("value_text") or "",
        num if num is not None else "",
        r.get("unit") or "",
        r["ref_low"] if r.get("ref_low") is not None else "",
        r["ref_high"] if r.get("ref_high") is not None else "",
        _flag_text(flag, lng),
        r.get("source_filename") or "",
        num if (flag == "high" and num is not None) else "",
        num if (flag == "low" and num is not None) else "",
    ])
    fill = ABNORMAL_HIGH_FILL if flag == "high" else (ABNORMAL_LOW_FILL if flag == "low" else None)
    if fill:
        for cell in ws[ws.max_row][:10]:
            cell.fill = fill


def _hide_aux_columns(ws) -> None:
    for idx in (_COL_HIGH_PT, _COL_LOW_PT):
        ws.column_dimensions[get_column_letter(idx)].hidden = True


def _add_trend_chart(ws, title: str, unit: str = "") -> None:
    """在数据右侧插入折线图：实测线 + 参考上下限虚线 + 异常点标记。

    使用 Excel 原生图表（非图片），打开后可直接编辑，且随数据联动。
    """
    n = ws.max_row
    if n < 2:
        return
    chart = LineChart()
    chart.title = title
    chart.height = 8.6
    chart.width = 21

    chart.add_data(Reference(ws, min_col=_COL_VALUE, min_row=1, max_row=n), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=_COL_REF_LOW, min_row=1, max_row=n), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=_COL_REF_HIGH, min_row=1, max_row=n), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=_COL_HIGH_PT, min_row=1, max_row=n), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=_COL_LOW_PT, min_row=1, max_row=n), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))

    # 实测线
    s = chart.series[0]
    s.graphicalProperties = GraphicalProperties()
    s.graphicalProperties.line.solidFill = _C_LINE
    s.graphicalProperties.line.width = 28575  # ≈2.25pt
    s.smooth = True
    s.marker = Marker(symbol="circle", size=5)
    s.marker.graphicalProperties = GraphicalProperties(solidFill=_C_LINE)
    s.marker.graphicalProperties.line.solidFill = _C_LINE

    # 参考上下限：虚线、无标记
    for idx in (1, 2):
        r = chart.series[idx]
        r.graphicalProperties = GraphicalProperties()
        r.graphicalProperties.line.solidFill = _C_REF
        r.graphicalProperties.line.dashStyle = "dash"
        r.graphicalProperties.line.width = 12700
        r.marker = Marker(symbol="none")

    # 异常点：只画标记、不连线
    for idx, color in ((3, _C_HIGH), (4, _C_LOW)):
        a = chart.series[idx]
        a.graphicalProperties = GraphicalProperties()
        a.graphicalProperties.line.noFill = True
        a.marker = Marker(symbol="circle", size=9)
        a.marker.graphicalProperties = GraphicalProperties(solidFill=color)
        a.marker.graphicalProperties.line.solidFill = color

    if unit:
        chart.y_axis.title = unit
    chart.legend.position = "b"
    chart.legend.overlay = False
    ws.add_chart(chart, _CHART_ANCHOR)


def _safe_sheet_name(name: str, used: set[str]) -> str:
    """生成合法且不重复的工作表名（Excel 限制 31 字符、禁用 []:*?/\\）。"""
    base = _INVALID_SHEET_CHARS.sub("", (name or "").strip())[:28] or "item"
    candidate, i = base, 1
    while candidate in used:
        i += 1
        candidate = f"{base[:24]}_{i}"
    used.add(candidate)
    return candidate


def export_items_xlsx(patient_id: int, items: list[str], lang: Optional[str] = None) -> bytes:
    """批量导出多个检验项目的历次序列：首张为汇总表，其余每个项目一张表。"""
    lng = lang or "zh-CN"
    wb = Workbook()
    summary = wb.active
    summary.title = tr(lng, "sheet.summary")
    headers = _trend_headers(lng)
    _write_header(summary, headers)
    used = {summary.title}

    for raw in items:
        norm = normalize_text(raw)
        series = dao.trend_series(patient_id, norm)
        if not series:
            continue
        for r in series:
            _append_trend_row(summary, r, lng)
        ws = wb.create_sheet(_safe_sheet_name(series[0].get("item") or raw, used))
        _write_header(ws, headers)
        for r in series:
            _append_trend_row(ws, r, lng)
        _sheet_widths(ws, headers)
        _style_body(ws)
        _hide_aux_columns(ws)
        unit = next((r.get("unit") for r in series if r.get("unit")), "") or ""
        _add_trend_chart(ws, f"{series[0].get('item') or raw} · {tr(lng, 'sheet.trend')}", unit)

    _sheet_widths(summary, headers)
    _style_body(summary)
    _hide_aux_columns(summary)
    return _to_bytes(wb)


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
