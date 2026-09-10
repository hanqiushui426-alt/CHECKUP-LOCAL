"""端到端 API 校验：患者 -> 指标 -> 趋势 -> Excel 导出 -> 模板。"""
from __future__ import annotations

import json
import urllib.parse
import urllib.request

BASE = "http://127.0.0.1:8321"


def get(path: str):
    with urllib.request.urlopen(BASE + path, timeout=10) as resp:
        return resp.status, resp.read()


def main() -> None:
    s, body = get("/api/stats/overview")
    ov = json.loads(body)
    assert ov["patients"] >= 1 and ov["results"] >= 1, ov
    pid = ov["patients"]
    print("overview ok:", pid, "patient(s),", ov["results"], "results")

    # 取第一位患者做明细校验（不假设具体 id=1）
    s, body = get("/api/patients?limit=10")
    first = json.loads(body)["items"][0]
    pid = first["id"]
    assert first["name"] and first["report_count"] >= 1, first
    print("patient ok:", first["name"], "reports:", first["report_count"])

    s, body = get("/api/stats/items?patient_id=%d" % pid)
    items = json.loads(body)
    assert items, items
    print("items ok:", items[0]["item"], "(计数项:", len(items), ")")

    # 定性指标（如"阴性"）没有数值，需挑一个可绘图的数值型指标
    item, trend, q = None, None, ""
    for cand in items:
        cq = urllib.parse.quote(cand["item"])
        s, body = get(f"/api/stats/trend?patient_id={pid}&item={cq}")
        t = json.loads(body)
        if t.get("series") and any(p.get("value_num") is not None for p in t["series"]):
            item, trend, q = cand["item"], t, cq
            break
    assert item and trend, "未找到可绘图的数值型指标"
    print("trend ok:", trend["meta"]["item"], "points:", len(trend["series"]))

    s, body = get(f"/api/patients/{pid}/reports?limit=500")
    reports = json.loads(body)
    assert reports["total"] >= 1
    rep = None
    for r in reports["items"]:  # 取一份带检验结果的报告（可能存在纯文字型报告 0 项）
        s, body = get(f"/api/reports/{r['id']}")
        cand = json.loads(body)
        if cand["patient_id"] == pid and (cand.get("patient") or {}).get("name") \
                and len(cand.get("results") or []) >= 1:
            rep = cand
            break
    assert rep, "未找到含检验结果的报告"
    print("report ok: results rows:", len(rep["results"]))
    total_results = sum(int(r["result_count"]) for r in reports["items"])
    print("patient total results:", total_results, "across", len(reports["items"]), "reports")

    # 导出 Excel（OpenPyXL 二次解析验证；行数 = 结果数 + 1 表头）
    s, body = get(f"/api/export/results.xlsx?patient_id={pid}")
    assert s == 200 and body[:2] == b"PK"
    from openpyxl import load_workbook
    from io import BytesIO

    wb = load_workbook(BytesIO(body))
    ws = wb["检验结果明细"]
    assert ws.max_row == total_results + 1 and ws.max_column == 14, (ws.max_row, total_results)
    print("export results.xlsx ok: sheet rows:", ws.max_row)

    s, body = get(f"/api/export/trend.xlsx?patient_id={pid}&item={q}")
    wb2 = load_workbook(BytesIO(body))
    print("export trend.xlsx ok: sheet:", wb2.sheetnames, "rows:", wb2.active.max_row)

    s, body = get("/api/templates")
    tpls = json.loads(body)
    assert len(tpls) >= 2
    print("templates ok:", [t["name"] for t in tpls])

    print("\nALL E2E CHECKS PASSED")


if __name__ == "__main__":
    main()
